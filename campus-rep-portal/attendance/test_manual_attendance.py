from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Department
from .models import AttendanceRecord, LectureSession


User = get_user_model()


class ManualAttendanceTests(TestCase):
    def setUp(self):
        self.department = Department.objects.create(
            name="Computer Science",
            faculty="Science",
        )
        self.rep = User.objects.create_user(
            username="rep",
            password="StrongPass123!",
            role="CLASS_REP",
            department=self.department,
            level="200",
            first_name="Class",
            last_name="Rep",
        )
        self.student = User.objects.create_user(
            username="student",
            password="StrongPass123!",
            role="STUDENT",
            department=self.department,
            level="200",
            first_name="Ada",
            last_name="Student",
            matric_number="CSC/2024/001",
        )
        self.absent_student = User.objects.create_user(
            username="absent",
            password="StrongPass123!",
            role="STUDENT",
            department=self.department,
            level="200",
            first_name="Grace",
            last_name="Absent",
            matric_number="CSC/2024/002",
        )
        self.session = LectureSession.objects.create(
            department=self.department,
            level="200",
            course_code="CSC201",
            venue_name="Lecture Hall",
            latitude=6.8,
            longitude=3.5,
            is_active=True,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.rep)

    def test_rep_can_mark_active_classmate_present_and_duplicate_is_safe(self):
        url = f"/api/attendance/classmates/{self.student.pk}/mark-attendance/"
        response = self.client.post(
            url,
            {"session_id": self.session.pk},
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertTrue(response.data["created"])
        record = AttendanceRecord.objects.get(
            student=self.student,
            session=self.session,
        )
        self.assertEqual(record.check_in_method, "REP_MANUAL")
        self.assertEqual(record.marked_by, self.rep)

        duplicate = self.client.post(
            url,
            {"session_id": self.session.pk},
            format="json",
        )
        self.assertEqual(duplicate.status_code, 200)
        self.assertFalse(duplicate.data["created"])
        self.assertEqual(
            AttendanceRecord.objects.filter(
                student=self.student,
                session=self.session,
            ).count(),
            1,
        )

    def test_rep_cannot_mark_student_outside_their_class_or_inactive_session(self):
        other_department = Department.objects.create(
            name="Mathematics",
            faculty="Science",
        )
        other_student = User.objects.create_user(
            username="other",
            password="StrongPass123!",
            role="STUDENT",
            department=other_department,
            level="200",
        )
        response = self.client.post(
            f"/api/attendance/classmates/{other_student.pk}/mark-attendance/",
            {"session_id": self.session.pk},
            format="json",
        )
        self.assertEqual(response.status_code, 404)

        self.session.is_active = False
        self.session.save(update_fields=["is_active"])
        inactive_response = self.client.post(
            f"/api/attendance/classmates/{self.student.pk}/mark-attendance/",
            {"session_id": self.session.pk},
            format="json",
        )
        self.assertEqual(inactive_response.status_code, 404)

    def test_students_cannot_manually_mark_attendance(self):
        self.client.force_authenticate(self.student)
        response = self.client.post(
            f"/api/attendance/classmates/{self.absent_student.pk}/mark-attendance/",
            {"session_id": self.session.pk},
            format="json",
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(
            AttendanceRecord.objects.filter(
                student=self.absent_student,
                session=self.session,
            ).exists()
        )

    def test_session_creation_succeeds_when_notifications_fail(self):
        payload = {
            "course_code": "CSC202",
            "venue_name": "Lecture Hall 2",
            "level": "200",
            "latitude": 6.8,
            "longitude": 3.5,
            "radius_meters": 50,
        }

        with patch(
            "attendance.views.Notification.objects.bulk_create",
            side_effect=RuntimeError("notification service unavailable"),
        ):
            response = self.client.post(
                "/api/attendance/sessions/create/",
                payload,
                format="json",
            )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(LectureSession.objects.filter(course_code="CSC202").exists())

    def test_rep_cannot_manage_a_session_for_another_level(self):
        other_level_session = LectureSession.objects.create(
            department=self.department,
            level="100",
            course_code="CSC101",
            venue_name="Lecture Hall 1",
            latitude=6.8,
            longitude=3.5,
        )

        response = self.client.post(
            f"/api/attendance/sessions/{other_level_session.pk}/toggle/",
            {},
            format="json",
        )

        self.assertEqual(response.status_code, 404)
        other_level_session.refresh_from_db()
        self.assertFalse(other_level_session.is_active)

    def test_super_admin_can_load_session_roster_and_mark_student_present(self):
        admin = User.objects.create_superuser(
            username="owner",
            password="test-password",
        )
        self.client.force_authenticate(admin)

        roster_response = self.client.get(
            f"/api/attendance/admin/sessions/{self.session.pk}/roster/"
        )
        self.assertEqual(roster_response.status_code, 200)
        self.assertEqual(
            {student["id"] for student in roster_response.data},
            {self.student.pk, self.absent_student.pk},
        )

        response = self.client.post(
            f"/api/attendance/classmates/{self.absent_student.pk}/mark-attendance/",
            {"session_id": self.session.pk},
            format="json",
        )
        self.assertEqual(response.status_code, 201)
        record = AttendanceRecord.objects.get(
            student=self.absent_student,
            session=self.session,
        )
        self.assertEqual(record.marked_by, admin)
        self.assertEqual(record.check_in_method, "REP_MANUAL")

    def test_export_lists_full_class_roster_with_presence_status(self):
        AttendanceRecord.objects.create(
            student=self.student,
            session=self.session,
            check_in_method="REP_MANUAL",
            marked_by=self.rep,
        )

        response = self.client.get(
            f"/api/attendance/sessions/{self.session.pk}/export/"
        )
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook["Attendance"]
        self.assertEqual(
            [sheet.cell(row=5, column=column).value for column in range(1, 7)],
            ["S/N", "Full Name", "Matric Number", "Attendance", "Time", "Marked by"],
        )
        rows = list(sheet.iter_rows(min_row=6, values_only=True))
        by_matric = {row[2]: row for row in rows}
        self.assertEqual(by_matric["CSC/2024/001"][3], "Present")
        self.assertEqual(by_matric["CSC/2024/001"][5], "Class Rep")
        self.assertEqual(by_matric["CSC/2024/002"][3], "Absent")
